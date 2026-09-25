"""Fill a copy of the supplied AI disclosure without inventing team sign-off."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "data" / "official" / "templates" / "LangAI3.0_AI_Disclosure.docx"
OUTPUT = ROOT / "submission" / "PRISM_Theme2_AI_Disclosure_Working.docx"

REPLACEMENTS = {
    3: "Team Name: To be supplied by the team",
    4: "Project / Product Name: PRISM Theme 2 Guided Troubleshooting Studio",
    5: "Organization / Institution: To be supplied by the team",
    6: "Submission Date: To be confirmed at submission",
    9: "Did your team use AI in developing this project? Yes. OpenAI Codex assisted with implementation and artifacts.",
    10: "Human review, acceptance, and representative sign-off are pending.",
    13: "Idea generation / brainstorming: Yes, architecture and evaluation planning.",
    14: "Code generation or assistance: Yes, Python API, compiler, caches, tests, and frontend code.",
    15: "UI / UX design: Yes, React console structure and copy.",
    16: "Content creation: Yes, draft presentation, disclosure, and demo script.",
    17: "Data analysis: Yes, local contract and benchmark report interpretation.",
    18: "Testing / debugging: Yes, fixtures, regressions, and release checks.",
    19: "Other: No real customer data, device control, or paid provider call was used in this workspace.",
    22: "1. Feature Name: Source-grounded compiler, API, catalog resolver, and cache",
    23: "2. Origin: AI-assisted implementation; team acceptance is pending",
    24: "3. Tool and prompt summary: OpenAI Codex followed the team's Theme 2 specification to draft code, tests, and fixes. Outputs were changed after deterministic checks. Independent semantic review remains pending.",
    26: "1. Feature Name: Preview console, evaluation fixtures, reports, and submission drafts",
    27: "2. Origin: AI-assisted implementation; team acceptance is pending",
    28: "3. Tool and prompt summary: OpenAI Codex drafted the interface, team-authored development fixtures, review sheets, and evidence reports. Human labels and variation review have not been supplied.",
    32: "AI usage complies with guidelines and policies: Team representative to confirm.",
    33: "No proprietary or copyrighted data misused: Team representative to confirm.",
    36: "Name of Team Representative: To be completed by the team",
    37: "Role: To be completed by the team",
    38: "Signature: To be completed by the team",
    39: "Date: To be completed by the team",
}


def main() -> int:
    document = Document(TEMPLATE)
    if len(document.paragraphs) < 40:
        raise ValueError("Official disclosure form structure changed")
    for index, value in REPLACEMENTS.items():
        document.paragraphs[index].text = value
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    with zipfile.ZipFile(OUTPUT) as archive:
        if archive.testzip() is not None:
            raise ValueError("Damaged DOCX archive")
    checked = Document(OUTPUT)
    if any(checked.paragraphs[index].text != value for index, value in REPLACEMENTS.items()):
        raise ValueError("Disclosure text changed during save")
    print(f"DISCLOSURE: {len(REPLACEMENTS)} form fields filled; {OUTPUT}")
    print("LIMIT: representative identity, attestation, signature and visual rendering remain pending")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
