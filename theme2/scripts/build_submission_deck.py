"""Fill the supplied 12-slide template with evidence-bounded working copy."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "data" / "official" / "templates" / "CollegeName_TeamName_Submission.pptx"
OUTPUT = ROOT / "submission" / "PRISM_Theme2_Submission_Working.pptx"

CONTENT: dict[int, list[str]] = {
    2: [
        "Theme 02: guided troubleshooting from the supplied SIIS article",
        "A complaint can point to a long article with several device topics.",
        "The API must return a complete, valid guide with approved catalog links.",
        "Source conditions must remain visible when customer facts are unknown.",
    ],
    3: [
        "A plausible source reference can still support the wrong operation.",
        "Partial catalog word overlap previously selected unrelated settings.",
        "A cached response can be unsafe when the operation or facts change.",
        "The prior hybrid remains a development baseline. A fair live comparison is pending.",
    ],
    4: [
        "1  Validate the complaint and immutable article payload.",
        "2  Check a versioned exact answer, then normalize explicit facts.",
        "3  Compile every actionable article section and record source offsets.",
        "4  Guard setting identity and operation before copying catalog metadata.",
        "5  Render the full official response. Project relevance separately for preview.",
    ],
    5: [
        "Mixed black-screen source: the official guide retains all source topics.",
        "The preview flags Kids and fingerprint sections as off-topic.",
        "Touch source: enable and disable remain separate catalog operations.",
        "Synthetic foldable case: inner and cover display facts stay distinct.",
        "Local demo: run the API, open the console, then inspect Evidence & trace.",
    ],
    6: [
        "API: FastAPI, Pydantic v2, one worker.",
        "Compilation: source blocks, coverage ledger, guarded catalog resolver.",
        "Retrieval: BM25; optional local MiniLM directory was not available here.",
        "State: bounded SQLite and memory caches with source/version identity.",
        "Console: React, TypeScript, Vite. Container build is provided but unverified.",
    ],
    7: [
        "Agents can inspect which source text supports each generated action.",
        "Unknown conditions remain visible for confirmation.",
        "A catalog link is an available action, not proof of device execution.",
        "The official response keeps the full article; the preview narrows the view.",
    ],
    8: [
        "Local test count and 20/20 supplied-query contract cases passed.",
        "3/3 team-authored unfamiliar articles passed contract checks.",
        "6/6 authored fact pairs and 30/30 cache adversaries passed local checks.",
        "Local loopback: 250/250 exact hits; 180/180 first generated variants hit.",
        "Independent semantic labels, hosted models and public latency are pending.",
        "Evidence: selected local release summary.",
    ],
    9: [
        "Two reviewers must label source actions, relevance and applicability.",
        "Obtain three sealed unfamiliar articles and six sealed fact pairs.",
        "Verify candidate model access, quota, fidelity and the cold deadline.",
        "Measure a public endpoint only after publication is authorized.",
        "Review variations, team details, disclosure and media before submission.",
    ],
    10: [
        "Full-source official output and a separate request-specific preview.",
        "Source offsets and a coverage ledger expose what was used.",
        "Setting and operation must match before a catalog URI is copied.",
        "Fact changes may reuse a procedure; answer reuse needs compatible facts.",
        "No official competition score or independent accuracy claim is made.",
    ],
    11: [
        "Working prototype code: local package ready; GitHub sharing pending.",
        "Reproducible README and clean source-copy check: yes.",
        "Official result export: 20 lines, contract validated locally.",
        "Presentation and AI disclosure: working copies in submission/.",
        "Demo video: working local copy; public endpoint awaits approval.",
    ],
}


def fill_lines(slide, lines: list[str], *, size: int = 23) -> None:
    frame = slide.shapes[1].text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.14)
    frame.margin_right = Inches(0.14)
    frame.margin_top = Inches(0.12)
    frame.margin_bottom = Inches(0.1)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(14)
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Calibri"
        run.font.size = Pt(size)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--release-run", required=True, help="Completed release report path relative to theme2")
    args = parser.parse_args()
    run = (ROOT / args.release_run).resolve()
    if ROOT not in run.parents or not run.is_dir():
        raise ValueError("Release report must be inside theme2")
    summary = json.loads((run / "summary.json").read_text(encoding="utf-8"))
    if summary["failed"] or summary["submission_export_lines"] != 20:
        raise ValueError("Release evidence is incomplete")
    test_match = re.search(r"(\d+) passed", (run / "tests.stdout.txt").read_text(encoding="utf-8"))
    if test_match is None:
        raise ValueError("Release test count is unavailable")
    count = int(test_match.group(1))
    CONTENT[8][0] = f"{count}/{count} local tests and 20/20 supplied-query contract cases passed."
    CONTENT[8][-1] = f"Evidence: reports/{run.name}/summary.json."
    presentation = Presentation(TEMPLATE)
    if len(presentation.slides) != 12:
        raise ValueError("The official submission template no longer has 12 slides")
    cover = presentation.slides[0].shapes[5]
    cover.height = Inches(2.45)
    cover.text = (
        "Theme ID - 02\n"
        "Team Name - to be supplied\n"
        "College Name - to be supplied\n"
        "Members and emails - to be supplied\n"
        "GitHub link - pending authorized sharing"
    )
    for paragraph in cover.text_frame.paragraphs:
        paragraph.space_after = Pt(4)
        for run in paragraph.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(16)
    for number, lines in CONTENT.items():
        fill_lines(presentation.slides[number - 1], lines, size=21 if number == 8 else 23)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    with zipfile.ZipFile(OUTPUT) as archive:
        if archive.testzip() is not None:
            raise ValueError("Damaged PPTX archive")
    checked = Presentation(OUTPUT)
    if len(checked.slides) != 12 or any(not checked.slides[number - 1].shapes[1].text.strip() for number in CONTENT):
        raise ValueError("Missing slide content")
    print(f"DECK: {len(checked.slides)} template slides filled; {OUTPUT}")
    print("LIMIT: team details, GitHub, visual render, independent labels and public metrics remain pending")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
